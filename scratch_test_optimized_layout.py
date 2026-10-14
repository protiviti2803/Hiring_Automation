import sqlite3
import json
import os
import math
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN

from app.services.pdf_export_service import (
    get_valid_projects,
    flatten_bullets,
    strip_bold,
    sanitize_bullet,
    clean_expertise_keyword,
    consolidate_versions,
    render_summary,
    render_bullet,
    build_standard_title,
    delete_paragraphs_after,
    apply_barlow_body_styling,
    get_key
)

# Realistic characters per line for 7.50" (540pt) text frame
BODY_CHARS_PER_LINE = 115
TITLE_CHARS_PER_LINE = 100
USABLE_HEIGHT_PT = 450.0  # Height of Text Placeholder 144 is 6.23" = 448.5 pt

def calc_lines(text: str, cpl: int) -> int:
    text = text.strip()
    if not text: return 0
    return max(1, math.ceil(len(text) / cpl))

def est_rendered_height(summary: str, projects: list, bpp_list: list, cfg: dict) -> float:
    # 1. Professional Experience Header
    total = 14.0
    # 2. Professional Summary
    sum_lines = calc_lines(strip_bold(sanitize_bullet(summary)), int(BODY_CHARS_PER_LINE * 8.0 / cfg["body_font_pt"]))
    total += sum_lines * cfg["summary_line_spacing"] + 3.0
    # 3. Major Projects Section Header
    total += 14.0 + 11.0 + 4.0
    
    # 4. Projects
    for i, proj in enumerate(projects):
        raw_title = strip_bold(proj.get('project_title') or proj.get('title') or '')
        clean_role = proj.get('role') or ''
        std_title = build_standard_title(raw_title, clean_role)
        t_cpl = int(TITLE_CHARS_PER_LINE * 8.5 / cfg["title_font_pt"])
        t_lines = calc_lines(std_title, t_cpl)
        total += cfg["space_before_title"] + t_lines * cfg["title_spacing_pt"] + cfg["space_after_title"]
        
        bpp = bpp_list[i] if i < len(bpp_list) else 2
        bullets = flatten_bullets(proj.get('bullets', []))
        b_cpl = int(BODY_CHARS_PER_LINE * 8.0 / cfg["body_font_pt"])
        for b in bullets[:bpp]:
            b_lines = calc_lines(b, b_cpl)
            total += 1.0 + b_lines * cfg["body_spacing_pt"] + cfg["space_after_bullet"]
            
    return total

def compute_optimized_layout(summary: str, projects: list) -> dict:
    H = USABLE_HEIGHT_PT
    n = len(projects)
    
    # Target bullets per project based on project count
    if n <= 2:
        ideal_bpp = 5
    elif n == 3:
        ideal_bpp = 4
    elif n <= 4:
        ideal_bpp = 3
    elif n <= 6:
        ideal_bpp = 2
    else:
        ideal_bpp = 2

    # Baseline configuration
    cfg = dict(
        title_font_pt        = 8.5,
        body_font_pt         = 8.0,
        title_spacing_pt     = 11.0,
        body_spacing_pt      = 10.5,
        space_before_title   = 4.0,
        space_after_title    = 1.0,
        space_after_bullet   = 1.5,
        summary_line_spacing = 11.0,
    )
    
    current_bpp = [
        min(ideal_bpp, max(1, len(flatten_bullets(p.get('bullets', [])))))
        for p in projects
    ]
    
    # If 4 projects, can we give 3 or 4 bullets? Let's check
    # If projects are few, try higher bullets first if they fit
    if n <= 4:
        for i in range(len(current_bpp)):
            avail = len(flatten_bullets(projects[i].get('bullets', [])))
            if avail > current_bpp[i]:
                # Test increasing
                test_bpp = list(current_bpp)
                test_bpp[i] += 1
                if est_rendered_height(summary, projects, test_bpp, cfg) <= H * 0.98:
                    current_bpp[i] += 1

    # Check if baseline fits
    h = est_rendered_height(summary, projects, current_bpp, cfg)
    
    # If overflowing, first adjust spacing/font slightly before stripping bullets
    if h > H:
        cfg.update(dict(
            title_font_pt        = 8.0,
            body_font_pt         = 7.5,
            title_spacing_pt     = 10.0,
            body_spacing_pt      = 9.5,
            space_before_title   = 3.0,
            space_after_title    = 0.5,
            space_after_bullet   = 1.0,
            summary_line_spacing = 10.0,
        ))
        h = est_rendered_height(summary, projects, current_bpp, cfg)
        
    # If still overflowing, reduce bullets
    if h > H:
        for _ in range(60):
            if est_rendered_height(summary, projects, current_bpp, cfg) <= H * 0.99:
                break
            reducible = [(i, b) for i, b in enumerate(current_bpp) if b > 1]
            if not reducible:
                break
            idx_reduce = max(reducible, key=lambda x: x[1])[0]
            current_bpp[idx_reduce] -= 1

    cfg["bullets_per_project_list"] = current_bpp
    final_h = est_rendered_height(summary, projects, current_bpp, cfg)
    print(f"[Optimized Layout] projects={n}, bpp={current_bpp} (Total {sum(current_bpp)} bullets), est_h={final_h:.1f}/{H}pt ({final_h/H*100:.1f}% fill), fonts={cfg['title_font_pt']}/{cfg['body_font_pt']}pt")
    return cfg

def render_optimized_pptx(data: dict, template_path: str, output_path: str):
    prs = Presentation(template_path)

    candidate_name = get_key(data, "candidate_name") or "Candidate Name"
    summary        = sanitize_bullet(get_key(data, "professional_summary") or "")

    # Sidebar data
    areas = get_key(data, "areas_of_expertise") or []
    if isinstance(areas, str): areas = [areas]
    areas_clean = consolidate_versions([clean_expertise_keyword(sanitize_bullet(i)) for i in areas if i])[:20]

    industry = get_key(data, "industry_expertise") or []
    if isinstance(industry, str): industry = [industry]
    industry_clean = [sanitize_bullet(i) for i in industry]

    education = get_key(data, "education") or get_key(data, "education_history") or ["Bachelor of Technology"]
    if isinstance(education, str): education = [education]
    education_clean = [sanitize_bullet(i) for i in education]

    raw_projects = get_key(data, "major_project_experience") or []
    projects = get_valid_projects(raw_projects)

    layout    = compute_optimized_layout(summary, projects)
    bpp_list  = layout["bullets_per_project_list"]
    title_fpt = layout["title_font_pt"]
    body_fpt  = layout["body_font_pt"]
    title_spt = layout["title_spacing_pt"]
    body_spt  = layout["body_spacing_pt"]
    sp_before = layout["space_before_title"]
    sp_after_t= layout["space_after_title"]
    sp_after_b= layout["space_after_bullet"]
    sum_spt   = layout["summary_line_spacing"]

    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame

            # Candidate name
            if "Title 138" in shape.name or any("FirstName" in p.text or "LastName" in p.text for p in tf.paragraphs):
                for p in tf.paragraphs:
                    if "FirstName" not in p.text and "LastName" not in p.text: continue
                    fn = fs = fi = fc = None
                    if p.runs and p.runs[0].font:
                        fn = p.runs[0].font.name
                        fs = p.runs[0].font.size
                        fi = p.runs[0].font.italic
                        try: fc = p.runs[0].font.color.rgb
                        except: pass
                    p.text = candidate_name.strip()
                    for r in p.runs:
                        if fn: r.font.name = fn
                        if fs: r.font.size = fs
                        r.font.bold = True
                        if fi is not None: r.font.italic = fi
                        if fc: r.font.color.rgb = fc

            # Sidebar
            elif "Placeholder 145" in shape.name or any(h in "".join(p.text for p in tf.paragraphs) for h in ["AREAS OF EXPERTISE", "INDUSTRY EXPERTISE", "EDUCATION"]):
                shape.top = Inches(0.2)
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(5)
                hstyles = []
                for hi in [0, 2, 4]:
                    s = {"level": 0, "font_name": None, "font_size": None, "font_bold": True, "font_color": None}
                    if hi < len(tf.paragraphs):
                        po = tf.paragraphs[hi]
                        s["level"] = po.level
                        if po.runs and po.runs[0].font:
                            ro = po.runs[0].font
                            s["font_name"], s["font_size"] = ro.name, ro.size
                            s["font_bold"] = ro.bold
                            try: s["font_color"] = ro.color.rgb
                            except: pass
                    hstyles.append(s)
                tf.clear()

                def _hdr(text, hs, first=False):
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    p.text, p.level, p.space_after = text, hs["level"], Pt(4)
                    if p.runs:
                        r = p.runs[0]
                        if hs["font_name"]: r.font.name = hs["font_name"]
                        if hs["font_size"]: r.font.size = hs["font_size"]
                        r.font.bold = hs["font_bold"] if hs["font_bold"] is not None else True
                        if hs["font_color"]: r.font.color.rgb = hs["font_color"]

                def _item(text, fn=None):
                    p = tf.add_paragraph()
                    p.text = f"• {sanitize_bullet(text)}"
                    p.level, p.space_after, p.line_spacing = 0, Pt(2), Pt(11)
                    p.alignment = PP_ALIGN.LEFT
                    if p.runs:
                        r = p.runs[0]
                        r.font.size, r.font.bold = Pt(8.0), False
                        if fn: r.font.name = fn

                fn0 = hstyles[0]["font_name"]
                _hdr("AREAS OF EXPERTISE", hstyles[0], first=True)
                tf.paragraphs[-1].space_before = Pt(0)
                for item in areas_clean: _item(item, fn0)

                _hdr("INDUSTRY EXPERTISE", hstyles[1])
                tf.paragraphs[-1].space_before = Pt(10)
                for item in industry_clean: _item(item, fn0)

                _hdr("EDUCATION", hstyles[2])
                tf.paragraphs[-1].space_before = Pt(10)
                for item in education_clean: _item(item, fn0)

            # Main Body
            elif "Placeholder 144" in shape.name or any(h in "".join(p.text for p in tf.paragraphs) for h in ["PROFESSIONAL EXPERIENCE", "MAJOR PROJECTS"]):
                shape.top = Inches(0.2)
                tf.word_wrap = True

                for idx, p in enumerate(tf.paragraphs):
                    if "PROFESSIONAL EXPERIENCE" in p.text and idx + 1 < len(tf.paragraphs):
                        ps = tf.paragraphs[idx + 1]
                        fn = ps.runs[0].font.name if ps.runs and ps.runs[0].font else None
                        render_summary(ps, summary, font_name=fn or "Barlow", font_size_pt=body_fpt, line_spacing_pt=sum_spt)
                        break

                mp_idx = next((i for i, p in enumerate(tf.paragraphs) if "MAJOR PROJECTS" in p.text), -1)
                if mp_idx != -1:
                    p_mp = tf.paragraphs[mp_idx]
                    p_mp.space_before, p_mp.space_after = Pt(10), Pt(3)
                    ff = p_mp.runs[0].font.name if p_mp.runs and p_mp.runs[0].font else None

                    delete_paragraphs_after(tf, mp_idx + 1)

                    for proj_i, proj in enumerate(projects):
                        raw_title = proj.get('project_title') or proj.get('title') or ''
                        raw_role  = proj.get('role') or ''
                        std_title = build_standard_title(raw_title, raw_role)

                        p_hdr = tf.add_paragraph()
                        p_hdr.text         = std_title
                        p_hdr.level        = 0
                        p_hdr.space_before = Pt(sp_before)
                        p_hdr.space_after  = Pt(sp_after_t)
                        p_hdr.line_spacing = Pt(title_spt)
                        if p_hdr.runs:
                            r = p_hdr.runs[0]
                            r.font.bold = True
                            r.font.size = Pt(title_fpt)
                            if ff: r.font.name = ff

                        bpp = bpp_list[proj_i] if proj_i < len(bpp_list) else 2
                        flat_buls = flatten_bullets(proj.get('bullets', []))
                        for b in flat_buls[:bpp]:
                            p_bul = tf.add_paragraph()
                            p_bul.space_before = Pt(1)
                            p_bul.space_after  = Pt(sp_after_b)
                            p_bul.line_spacing = Pt(body_spt)
                            render_bullet(p_bul, b, font_name=ff or "Barlow", font_size_pt=body_fpt)

    apply_barlow_body_styling(prs)
    prs.save(output_path)
    print(f"Saved optimized presentation to {output_path}")

# Run tests
if __name__ == "__main__":
    conn = sqlite3.connect('resume_analyzer.db')
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    
    template_path = os.path.abspath("app/assests/Protiviti_Format_Resume.pptx")
    
    for jid in [68, 69]:
        cur.execute('SELECT id, candidate_name, result_json FROM cv_jobs WHERE id=?', (jid,))
        j = cur.fetchone()
        data = json.loads(j['result_json'])
        out_pptx = f"app/temp/test_optimized_job_{jid}.pptx"
        render_optimized_pptx(data, template_path, out_pptx)
    conn.close()
