import os
import openpyxl
import json

def analyze_excel():
    excel_path = os.path.join("app", "assests", "Demands_Questions.xlsx")
    wb = openpyxl.load_workbook(excel_path)
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n================ SHEET: {sheet_name} ================")
        headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
        print("Headers:", headers[:10])
        
        # Collect distinct roles or candidate titles in column 2 (index 1) or column 1
        roles = set()
        candidates = []
        for r in range(2, sheet.max_row + 1):
            cand_name = sheet.cell(r, 1).value
            role_val = sheet.cell(r, 2).value
            q_val = sheet.cell(r, 3).value
            if role_val:
                roles.add(str(role_val).strip())
            if cand_name or role_val or q_val:
                candidates.append({
                    "row": r,
                    "candidate": cand_name,
                    "role": role_val,
                    "has_questions": bool(q_val)
                })
        print(f"Total entries: {len(candidates)}")
        print("Distinct Roles found:", list(roles)[:15])

if __name__ == "__main__":
    analyze_excel()
