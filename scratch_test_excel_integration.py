import os
import openpyxl
import re

def load_excel_questions_for_role(role_name: str) -> list:
    excel_path = os.path.join("app", "assests", "Demands_Questions.xlsx")
    if not os.path.exists(excel_path):
        print("Excel not found")
        return []
    
    try:
        wb = openpyxl.load_workbook(excel_path)
        if "Interview Question" not in wb.sheetnames:
            return []
        sheet = wb["Interview Question"]
        
        role_questions = []
        role_name_clean = role_name.strip().lower()
        
        is_ba = "ba" in role_name_clean or "analyst" in role_name_clean
        is_pmo = "pmo" in role_name_clean
        is_pm = "pm" in role_name_clean and not is_pmo
        
        for r in range(2, sheet.max_row + 1):
            sheet_role = str(sheet.cell(r, 2).value or "").strip().lower()
            q_raw = sheet.cell(r, 3).value
            if not sheet_role or not q_raw:
                continue
                
            matched = False
            if role_name_clean in sheet_role or sheet_role in role_name_clean:
                matched = True
            elif is_pmo and "pmo" in sheet_role:
                matched = True
            elif is_ba and ("ba" in sheet_role or "analyst" in sheet_role):
                matched = True
            elif is_pm and "pm" in sheet_role:
                matched = True
                
            if matched:
                lines = [l.strip() for l in str(q_raw).split("\n") if l.strip()]
                for l in lines:
                    cleaned_q = re.sub(r'^\d+[\.\)\s]*', '', l).strip()
                    if len(cleaned_q) > 8 and not cleaned_q.lower().startswith("panelist") and not cleaned_q.lower().startswith("role"):
                        role_questions.append({
                            "question_text": cleaned_q,
                            "category": "Domain" if ("understanding" in cleaned_q.lower() or "role" in cleaned_q.lower()) else "Technical",
                            "source": "EXCEL_BASELINE"
                        })
                        
        return role_questions
    except Exception as e:
        print(f"Error: {e}")
        return []

if __name__ == "__main__":
    for test_role in ["Business Analyst (BA)", "PMO", "Asset Management PM", "Backend Engineer"]:
        qs = load_excel_questions_for_role(test_role)
        print(f"\nRole: [{test_role}] -> Extracted {len(qs)} Excel questions")
        for q in qs[:3]:
            print(f"  - [{q['source']}] {q['question_text'][:80]}...")
