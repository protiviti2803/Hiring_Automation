import os
from pptx import Presentation

prs = Presentation('app/assests/Protiviti_Format_Resume.pptx')
slide = prs.slides[0]
print(f"Slide width: {prs.slide_width/914400:.3f}\", height: {prs.slide_height/914400:.3f}\"")
for i, s in enumerate(slide.shapes):
    print(f"Shape {i}: '{s.name}', top={s.top/914400:.3f}\", left={s.left/914400:.3f}\", width={s.width/914400:.3f}\", height={s.height/914400:.3f}\"")
    if s.has_text_frame:
        print(f"  Paragraphs: {len(s.text_frame.paragraphs)}")
        for pi, p in enumerate(s.text_frame.paragraphs):
            print(f"    p{pi}: {repr(p.text[:70])}")
