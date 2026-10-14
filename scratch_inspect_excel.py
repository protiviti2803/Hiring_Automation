import os
import sys

def inspect_excel():
    excel_path = os.path.join("app", "assests", "Demands_Questions.xlsx")
    if not os.path.exists(excel_path):
        print(f"File not found: {excel_path}")
        return
        
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path)
        print("Sheet names:", wb.sheetnames)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            print(f"\n--- Sheet: {sheet_name} (Rows: {sheet.max_row}, Cols: {sheet.max_column}) ---")
            for r in range(1, min(6, sheet.max_row + 1)):
                row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
                print(f"Row {r}: {row_vals}")
    except ImportError:
        print("openpyxl is not installed.")

if __name__ == "__main__":
    inspect_excel()
