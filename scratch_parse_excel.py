import os
import openpyxl

def parse_questions_sheet():
    excel_path = os.path.join("app", "assests", "Demands_Questions.xlsx")
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb["Interview Question"]
    
    role_questions_map = {}
    
    for r in range(2, sheet.max_row + 1):
        panelist = sheet.cell(r, 1).value
        role = sheet.cell(r, 2).value
        questions_raw = sheet.cell(r, 3).value
        
        if role and questions_raw:
            role_str = str(role).strip()
            q_lines = [q.strip() for q in str(questions_raw).split("\n") if q.strip()]
            
            if role_str not in role_questions_map:
                role_questions_map[role_str] = []
                
            for line in q_lines:
                # Clean numbered prefix like "1.", "1Can", etc.
                cleaned_q = line
                import re
                cleaned_q = re.sub(r'^\d+[\.\)\s]*', '', cleaned_q).strip()
                if len(cleaned_q) > 5:
                    role_questions_map[role_str].append(cleaned_q)
                    
    print(f"Total roles mapped: {len(role_questions_map)}")
    for role, q_list in role_questions_map.items():
        print(f"\nRole: [{role}] -> {len(q_list)} questions extracted")
        for q in q_list[:3]:
            print(f"  - {q}")

if __name__ == "__main__":
    parse_questions_sheet()
