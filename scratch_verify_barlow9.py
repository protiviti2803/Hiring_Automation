import os
import json
import sqlite3
import subprocess
from pptx import Presentation
from app.services.pdf_export_service import render_pptx_template

template_path = os.path.abspath("app/assests/Protiviti_Format_Resume.pptx")

# 1. Test existing candidates from database
conn = sqlite3.connect('resume_analyzer.db')
conn.row_factory = sqlite3.Row
cur = conn.cursor()
cur.execute('SELECT id, candidate_name, result_json FROM cv_jobs WHERE status="COMPLETED"')
jobs = cur.fetchall()

print("=== 1. Testing Existing Profiles with Barlow 9pt ===")
for j in jobs:
    jid = j['id']
    cand_name = j['candidate_name'].encode('ascii', errors='replace').decode('ascii')
    print(f"\n--- Job {jid}: {cand_name} ---")
    data = json.loads(j['result_json'])
    out_pptx = os.path.abspath(f"app/temp/barlow9_job_{jid}.pptx")
    out_pdf = os.path.abspath(f"app/temp/barlow9_job_{jid}.pdf")
    render_pptx_template(data, template_path, out_pptx)
    
    # Check font sizes in output pptx
    prs = Presentation(out_pptx)
    slide = prs.slides[0]
    for shape in slide.shapes:
        if shape.has_text_frame:
            if "MAJOR PROJECTS" in shape.text_frame.text:
                for pi, p in enumerate(shape.text_frame.paragraphs):
                    if p.runs and p.runs[0].font.size:
                        font_size_pt = p.runs[0].font.size.pt
                        font_name = p.runs[0].font.name
                        safe_text = p.text[:60].encode('ascii', errors='replace').decode('ascii')
                        print(f"  P{pi:02d} ({font_name} {font_size_pt}pt): {safe_text}...")
                        
    # Convert to PDF
    subprocess.run(["python", "app/services/convert_pptx.py", out_pptx, out_pdf], capture_output=True)
    print(f"  Converted to: {out_pdf}")

conn.close()

# 2. Test Synthetic Profiles with FEW projects (1 project and 2 projects)
print("\n=== 2. Testing Profiles with Few Projects (1 & 2 projects) ===")

# Few projects test case 1: 2 Projects
test_2_projects = {
    "candidate_name": "Ananya Sharma",
    "professional_summary": "**Ananya Sharma** has **8 years** of experience as a **Senior Business Analyst** specialising in **Digital Banking Platforms**, **Payment Gateways**, and **Core Banking Transformation** within the **Banking and Financial Services** industry. She has led end-to-end SDLC initiatives for omni-channel customer journeys, open banking APIs, and high-volume transaction processing systems. Ananya is proficient in Agile Scrum ceremonies, BRD/FRD authoring, SQL data validation, JIRA workflow design, and user acceptance testing across multi-tier banking architectures.",
    "areas_of_expertise": ["Digital Banking", "Core Banking", "API Banking", "Payment Gateways", "SQL", "JIRA", "Confluence", "UAT", "Agile", "Scrum", "Business Process Modeling", "Finacle-10/11", "ISO 20022", "REST APIs", "Swagger", "Postman", "Defect Management", "Stakeholder Management", "Gap Analysis", "Wireframing"],
    "industry_expertise": ["Digital Banking", "Payment Systems", "Retail Banking"],
    "education": ["Master of Business Administration (MBA)", "Bachelor of Technology (Computer Science)"],
    "major_project_experience": [
        {
            "project_title": "Private Sector Bank – Omni-Channel Digital Banking Transformation",
            "role": "Lead Business Analyst",
            "bullets": [
                "Led the comprehensive requirement elicitation, process flow engineering, and functional architecture design for an enterprise digital banking platform serving over 4 million active retail and corporate customers.",
                "Authored 25+ detailed Business Requirement Documents (BRDs), Functional Specification Documents (FSDs), and API mapping matrices integrating core banking systems with modern microservices architectures.",
                "Engineered automated KYC and instant onboarding workflows leveraging India Stack components (Aadhaar e-KYC, DigiLocker, eSign, and Account Aggregator), slashing onboarding turnaround time from 2 days to under 6 minutes.",
                "Collaborated closely with enterprise security, fraud risk, compliance, and regulatory audit teams to ensure strict adherence to RBI cybersecurity and data localization mandates.",
                "Managed end-to-end User Acceptance Testing (UAT) across 15 sprint cycles, defining 200+ test scenarios, resolving critical defects in JIRA, and ensuring 99.98% sprint release stability."
            ]
        },
        {
            "project_title": "Payment Technology Services – Real-Time Payment Gateway & ISO 20022 Migration",
            "role": "Senior Consultant – Payments",
            "bullets": [
                "Spearheaded the business architecture and migration strategy for a high-throughput payment gateway processing over 12 million daily transactions across UPI, IMPS, NEFT, and RTGS payment rails.",
                "Designed end-to-end message mapping logic and data validation pipelines to transition legacy payment messaging formats to modern ISO 20022 XML standards.",
                "Conducted extensive data modeling and transaction reconciliation analyses using SQL, identifying settlement bottlenecks and reducing reconciliation discrepancies by 85%.",
                "Facilitated daily standups, backlog refinement, and sprint planning sessions as Proxy Product Owner, bridging communication between business stakeholders and distributed engineering teams.",
                "Delivered comprehensive functional training, user manuals, and operational SOPs to operations and customer support divisions, ensuring flawless go-live adoption."
            ]
        }
    ]
}

out_2p_pptx = os.path.abspath("app/temp/barlow9_2_projects.pptx")
out_2p_pdf = os.path.abspath("app/temp/barlow9_2_projects.pdf")
render_pptx_template(test_2_projects, template_path, out_2p_pptx)
subprocess.run(["python", "app/services/convert_pptx.py", out_2p_pptx, out_2p_pdf], capture_output=True)
print(f"2-Project Profile Saved & Converted to: {out_2p_pdf}")

# Few projects test case 2: 1 Project (Deep single project)
test_1_project = {
    "candidate_name": "Vikramaditya Roy",
    "professional_summary": "**Vikramaditya Roy** has **7 years** of experience as a **Senior Functional Consultant** specialising in **Enterprise Credit Risk Automation**, **Basel III/3.1 Capital Calculations**, and **Regulatory Reporting** within the **Investment Banking** industry. He specializes in designing robust risk analytics platforms, translating complex prudential regulations into automated calculation engines, and optimizing credit decisioning workflows. Vikramaditya collaborates closely with quantitative modelers, credit risk officers, and IT developers to streamline risk-weighted asset calculations, stress testing data pipelines, and regulatory submissions across global jurisdictions.",
    "areas_of_expertise": ["Credit Risk", "Basel III/3.1", "IFRS 9", "RWA Modeling", "Regulatory Reporting", "SQL", "Python", "Tableau", "JIRA", "UAT", "Agile", "PRA", "ECB", "Stress Testing", "CCAR", "ICAAP", "Market Risk", "Counterparty Risk", "Data Governance", "Risk Analytics"],
    "industry_expertise": ["Investment Banking", "Regulatory Risk", "Capital Markets"],
    "education": ["Master of Science in Financial Engineering", "Bachelor of Engineering"],
    "major_project_experience": [
        {
            "project_title": "Global Investment Bank – Enterprise Regulatory Capital & Stress Testing Engine",
            "role": "Lead Functional Risk Consultant",
            "bullets": [
                "Spearheaded the functional design and implementation of a multi-jurisdictional Basel III/3.1 Regulatory Capital calculation platform, automating RWA calculations for over $45B in wholesale exposures.",
                "Authored end-to-end functional specifications translating PRA and ECB capital rules, credit risk mitigation (CRM) haircut algorithms, and supervisory slotting criteria into scalable calculation logic.",
                "Engineered comprehensive data validation, reconciliation, and exception-handling frameworks using SQL and Python, ensuring 100% data integrity between front-office trade capture systems and regulatory risk datamarts.",
                "Led the functional delivery of automated CCAR and ICAAP stress-testing scenario engines, enabling risk officers to model macroeconomic shock impacts across credit portfolios in real-time.",
                "Designed and executed over 350 end-to-end UAT test cases covering complex derivative collateralizations, securitization tranches, and wholesale lending defaults, achieving zero post-deployment regulatory defects.",
                "Facilitated executive walkthroughs and steering committee presentations with Chief Risk Officers, regulators, and external auditors, establishing complete transparency and auditability of calculation methodologies."
            ]
        }
    ]
}

out_1p_pptx = os.path.abspath("app/temp/barlow9_1_project.pptx")
out_1p_pdf = os.path.abspath("app/temp/barlow9_1_project.pdf")
render_pptx_template(test_1_project, template_path, out_1p_pptx)
subprocess.run(["python", "app/services/convert_pptx.py", out_1p_pptx, out_1p_pdf], capture_output=True)
print(f"1-Project Profile Saved & Converted to: {out_1p_pdf}")
