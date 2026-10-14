import sqlite3
import json
import os
import sys

from app.services.pdf_export_service import (
    render_pptx_template,
    get_valid_projects,
    compute_layout,
    estimate_height_pt,
    USABLE_HEIGHT_PT,
    CHARS_PER_LINE
)
from pptx import Presentation
from pptx.util import Inches, Pt

conn = sqlite3.connect('resume_analyzer.db')
conn.row_factory = sqlite3.Row
cur = conn.cursor()
cur.execute('SELECT id, candidate_name, status, result_json, raw_text FROM cv_jobs WHERE status="COMPLETED" ORDER BY id DESC LIMIT 5')
jobs = cur.fetchall()

template_path = os.path.abspath("app/assests/Protiviti_Format_Resume.pptx")
print("Template path exists:", os.path.exists(template_path))

for j in jobs:
    print(f"\n==========================================")
    print(f"JOB ID {j['id']} - Candidate: {j['candidate_name']}")
    print(f"==========================================")
    data = json.loads(j['result_json'])
    
    summary = data.get('professional_summary', '')
    projects = data.get('major_project_experience', [])
    print(f"Summary words: {len(summary.split())}, chars: {len(summary)}")
    print(f"Total projects in JSON: {len(projects)}")
    
    valid_projects = get_valid_projects(projects)
    print(f"Valid projects after filtering: {len(valid_projects)}")
    
    layout = compute_layout(summary, valid_projects)
    bpp_list = layout["bullets_per_project_list"]
    print(f"Bullets per project allocated: {bpp_list}")
    
    # Test rendering PPTX
    out_pptx = f"app/temp/test_job_{j['id']}.pptx"
    try:
        render_pptx_template(data, template_path, out_pptx)
        print(f"Rendered PPTX to: {out_pptx}")
        
        # Now let's inspect the generated presentation shapes and paragraphs
        prs = Presentation(out_pptx)
        slide = prs.slides[0]
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text
                if "MAJOR PROJECTS" in txt or "PROFESSIONAL EXPERIENCE" in txt:
                    print(f"\n--- Shape: {shape.name} (Top: {shape.top/914400:.2f}\", Height: {shape.height/914400:.2f}\", Width: {shape.width/914400:.2f}\") ---")
                    print(f"Total paragraphs: {len(shape.text_frame.paragraphs)}")
                    # Let's list the paragraphs
                    for pi, p in enumerate(shape.text_frame.paragraphs):
                        runs_info = [(r.text[:30], r.font.size.pt if r.font.size else None, r.font.bold) for r in p.runs]
                        print(f"  P{pi:02d} (level={p.level}, sp_b={p.space_before.pt if p.space_before else None}, sp_a={p.space_after.pt if p.space_after else None}, linesp={p.line_spacing.pt if p.line_spacing else None}): {p.text[:70]}... | Runs: {len(p.runs)}")
    except Exception as e:
        import traceback
        print(f"Error rendering: {e}")
        traceback.print_exc()

conn.close()
